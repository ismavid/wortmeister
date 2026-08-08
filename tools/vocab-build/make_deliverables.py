#!/usr/bin/env python3
"""Emit the three step-1 deliverables: master XLSX, app seed JSON, Anki TSV."""
import json
import os
import re
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = os.environ.get("OUTDIR", "/sessions/gallant-stoic-faraday/mnt/_Agent/vocab")
os.makedirs(OUT, exist_ok=True)

vocab = json.load(open("/tmp/de/vocab.json", encoding="utf-8"))
vprep = json.load(open("/tmp/de/verb_prep.json", encoding="utf-8"))

# ---- attach verb rection to the matching verb rows
by_base = defaultdict(list)
for r in vprep:
    by_base[r["base"]].append(r)
for row in vocab:
    if row["pos"] == "verb":
        pats = by_base.get(row["lemma"], [])
        row["rection"] = " / ".join(p["pattern"] for p in pats)
    else:
        row["rection"] = ""

POS_DE = {"noun": "Nomen", "verb": "Verb", "adj": "Adjektiv", "adv": "Adverb",
          "conj": "Konjunktion", "prep": "Präposition", "pron": "Pronomen",
          "num": "Numerale", "det": "Artikelwort", "particle": "Partikel",
          "intj": "Interjektion", "prefix": "Präfix", "postp": "Postposition"}


def display(r):
    """Headword as a learner should memorise it: 'der Betrieb, -e'."""
    if r["pos"] == "noun":
        s = f'{r["article"]} {r["lemma"]}'.strip()
        return f'{s}, {r["plural"]}' if r["plural"] else s
    return r["lemma"]


def principal_parts(r):
    if r["pos"] != "verb":
        return ""
    bits = [r["lemma"], r["praeteritum"], f'{r["aux"]} {r["partizip2"]}'.strip()]
    return ", ".join(b for b in bits if b.strip())


COLS = [
    ("Level", 8), ("Tier", 12), ("German", 30), ("Article", 8), ("Lemma", 24),
    ("Plural", 20), ("POS", 13), ("English", 44), ("Präteritum", 18),
    ("Partizip II", 20), ("Aux", 8), ("3. Pers. Sg.", 16), ("Separable", 10),
    ("Verb + Präposition", 34), ("Other senses", 52), ("Freq. class", 11),
    ("Freq. rank", 11), ("Source", 20), ("Review", 8),
]

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=10)
LEVEL_FILL = {"A1": "E8F1DE", "A1*": "E8F1DE", "A2": "DDEBF7", "A2*": "DDEBF7",
              "B1": "FFF2CC", "B2": "FCE4D6"}
THIN = Side(style="thin", color="BFBFBF")


def write_sheet(ws, rows, freeze="A2"):
    ws.append([c[0] for c in COLS])
    for i, (name, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    for r in rows:
        ws.append([
            r["level"], r["tier"], display(r), r["article"], r["lemma"],
            r["plural"], POS_DE.get(r["pos"], r["pos"]), r["en"],
            r["praeteritum"], r["partizip2"], r["aux"], r["praesens_3sg"],
            r["separable"], r["rection"], r["en_senses"] or "",
            r["freq_class"], r["freq_rank"], r["source"],
            "check" if r["review"] else "",
        ])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(COLS)):
        lvl = row[0].value
        fill = PatternFill("solid", fgColor=LEVEL_FILL.get(lvl, "FFFFFF"))
        for c in row:
            c.font = BODY
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(vertical="top", wrap_text=False)
        row[0].fill = fill
        row[2].font = Font(name="Arial", size=10, bold=True)
    ws.freeze_panes = freeze
    if ws.max_row > 1:
        ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"
        t = Table(displayName=re.sub(r"\W", "", ws.title) + "Tbl", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1",
                                          showRowStripes=False)
        ws.add_table(t)


wb = Workbook()

# ---------- README
ws = wb.active
ws.title = "README"
readme = [
    ("German A1–B2 Vocabulary Master List", True),
    ("", False),
    ("Built for Goethe-Zertifikat B2, November 2026.", False),
    ("", False),
    ("HOW THIS WAS BUILT", True),
    ("A1 / A2 / B1 levels come from the official Goethe-Institut Wortlisten "
     "(the only word lists Goethe publishes).", False),
    ("Goethe publishes NO official B2 word list. The B2 layer is therefore "
     "constructed: it is every lemma in DeReWo (the IDS Mannheim corpus "
     "lemma-frequency list) down to frequency class 14 that is not already "
     "in the A1–B1 lists.", False),
    ("Gender and plural come from german-nouns (German Wiktionary). "
     "Verb forms, auxiliaries and English senses come from the English "
     "Wiktionary extract (kaikki.org). Verb + preposition patterns are "
     "hand-curated.", False),
    ("", False),
    ("COLUMNS", True),
    ("Tier — B2-core (freq class ≤13) is the priority study set; "
     "B2-extended (class 14) is the long tail. Study B2-core first.", False),
    ("Freq. class — DeReWo logarithmic frequency class. Lower = more common. "
     "Class 8 ≈ top 300 words, class 14 ≈ top 11,000.", False),
    ("Review — 'check' marks an entry whose gloss or gender came from a "
     "weaker source and is worth verifying before you drill it.", False),
    ("", False),
    ("KNOWN LIMITS", True),
    ("The B2 layer is frequency-derived from a news corpus, so it leans "
     "slightly journalistic and will contain some words no B2 exam needs, "
     "while a handful of B2 words below class 14 are missing.", False),
    ("Levels are cumulative: a word tagged B1 is expected knowledge at B2.", False),
    ("A1* / A2* mark very frequent function words (alle, am, zwei, uns) that "
     "the Goethe alphabetical lists omit. Their level is inferred from "
     "frequency, not taken from Goethe. Frequent CONTENT words are deliberately "
     "left at B2 - 'zusätzlich' is common in a news corpus but is not A1.", False),
]
for text, bold in readme:
    ws.append([text])
    ws.cell(row=ws.max_row, column=1).font = Font(
        name="Arial", size=12 if bold else 10, bold=bold)
ws.column_dimensions["A"].width = 118
for r in range(1, ws.max_row + 1):
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- data sheets
write_sheet(wb.create_sheet("ALL (A1-B2)"), vocab)
for lvl in ["A1", "A2", "B1", "B2"]:
    # A1* / A2* (frequency-inferred function words) live on their base sheet;
    # '*' is not a legal character in an Excel sheet name.
    write_sheet(wb.create_sheet(lvl),
                [r for r in vocab if r["level"].rstrip("*") == lvl])
write_sheet(wb.create_sheet("Nouns"), [r for r in vocab if r["pos"] == "noun"])
write_sheet(wb.create_sheet("Verbs"), [r for r in vocab if r["pos"] == "verb"])
write_sheet(wb.create_sheet("Adjectives+Adverbs"),
            [r for r in vocab if r["pos"] in ("adj", "adv")])
write_sheet(wb.create_sheet("Function words"),
            [r for r in vocab if r["pos"] in ("conj", "prep", "pron", "num",
                                              "det", "particle", "intj", "prefix")])

# ---------- verb + preposition sheet
vp = wb.create_sheet("Verb+Präposition")
vp.append(["Verb", "Präposition", "Kasus", "Muster", "English", "Beispiel", "Reflexiv"])
for i, w in enumerate([26, 14, 12, 34, 30, 48, 10], start=1):
    vp.column_dimensions[get_column_letter(i)].width = w
    c = vp.cell(row=1, column=i)
    c.fill, c.font = HDR_FILL, HDR_FONT
for r in vprep:
    vp.append([r["verb"], r["prep"], r["case"], r["pattern"], r["en"],
               r["example"], r["reflexive"]])
for row in vp.iter_rows(min_row=2, max_row=vp.max_row, max_col=7):
    for c in row:
        c.font = BODY
        c.border = Border(bottom=THIN)
    row[3].font = Font(name="Arial", size=10, bold=True)
vp.freeze_panes = "A2"

# ---------- stats
st = wb.create_sheet("Stats")
st.append(["Metric", "Value"])
for c in st[1]:
    c.fill, c.font = HDR_FILL, HDR_FONT
st.column_dimensions["A"].width = 34
st.column_dimensions["B"].width = 14
lv, ps = Counter(r["level"] for r in vocab), Counter(r["pos"] for r in vocab)
stats = [("Total entries", len(vocab))]
stats += [(f"Level {k}", lv[k]) for k in ["A1", "A1*", "A2", "A2*", "B1", "B2"]]
stats += [("Tier B2-core", sum(1 for r in vocab if r["tier"] == "B2-core")),
          ("Tier B2-extended", sum(1 for r in vocab if r["tier"] == "B2-extended"))]
stats += [(f"POS {POS_DE.get(k, k)}", v) for k, v in ps.most_common()]
stats += [("Nouns with article", sum(1 for r in vocab if r["pos"] == "noun" and r["article"])),
          ("Nouns with plural", sum(1 for r in vocab if r["pos"] == "noun" and r["plural"])),
          ("Verbs with Partizip II", sum(1 for r in vocab if r["pos"] == "verb" and r["partizip2"])),
          ("Separable verbs", sum(1 for r in vocab if r["separable"] == "yes")),
          ("Verbs taking sein", sum(1 for r in vocab if r["aux"] == "sein")),
          ("Verb+Präposition patterns", len(vprep)),
          ("Entries flagged for review", sum(1 for r in vocab if r["review"]))]
for k, v in stats:
    st.append([k, v])
for row in st.iter_rows(min_row=2, max_row=st.max_row, max_col=2):
    for c in row:
        c.font = BODY

xlsx_path = f"{OUT}/German_A1-B2_Vocabulary_Master.xlsx"
wb.save(xlsx_path)

# ---------- app seed JSON
app = [{
    "id": i + 1, "lemma": r["lemma"], "display": display(r), "pos": r["pos"],
    "level": r["level"], "tier": r["tier"], "article": r["article"] or None,
    "plural": r["plural"] or None, "en": r["en"],
    "senses": r["en_senses"], "principalParts": principal_parts(r) or None,
    "praeteritum": r["praeteritum"] or None, "partizip2": r["partizip2"] or None,
    "aux": r["aux"] or None, "praesens3sg": r["praesens_3sg"] or None,
    "separable": r["separable"] == "yes" if r["pos"] == "verb" else None,
    "rection": r["rection"] or None,
    "freqClass": r["freq_class"], "freqRank": r["freq_rank"],
    "source": r["source"], "needsReview": bool(r["review"]),
} for i, r in enumerate(vocab)]
json.dump({"meta": {"title": "German A1-B2 Vocabulary",
                    "built": "2026-08-08", "count": len(app),
                    "levels": dict(lv), "verbPrepositions": len(vprep)},
           "words": app, "verbPrepositions": vprep},
          open(f"{OUT}/german_vocab_a1_b2.json", "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)

# ---------- Anki TSV (Front, Back, Tags)
with open(f"{OUT}/anki_german_a1_b2.tsv", "w", encoding="utf-8") as fh:
    fh.write("#separator:tab\n#html:true\n#tags column:3\n")
    for r in vocab:
        front = display(r)
        back = r["en"]
        extra = []
        if r["pos"] == "verb":
            pp = principal_parts(r)
            if pp:
                extra.append(pp)
            if r["rection"]:
                extra.append(r["rection"])
        if extra:
            back += "<br><i>" + "<br>".join(extra) + "</i>"
        tags = f'{r["level"]} {POS_DE.get(r["pos"], r["pos"])} {r["tier"]}'.replace("-", "_")
        fh.write(f"{front}\t{back}\t{tags}\n")

with open(f"{OUT}/anki_verb_praeposition.tsv", "w", encoding="utf-8") as fh:
    fh.write("#separator:tab\n#html:true\n#tags column:3\n")
    for r in vprep:
        fh.write(f'{r["verb"]} + ___ ?\t{r["pattern"]}<br>{r["en"]}'
                 f'<br><i>{r["example"]}</i>\tVerbPraeposition\n')

print("wrote:")
for f in sorted(os.listdir(OUT)):
    print(f"  {OUT}/{f}  {os.path.getsize(os.path.join(OUT, f)):,} bytes")
